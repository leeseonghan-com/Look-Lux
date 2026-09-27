"""기능별 Excel 내보내기 (통합)

각 메뉴에서 "📥 엑셀 내보내기" 버튼으로 현재 보고 있는 데이터를 받을 수 있다.
  /export/projects   프로젝트
  /export/quotes     견적서 (+ 품목)
  /export/expenses   하청/외주 지급
  /export/items      견적 단가표
  /export/inventory  재고/장비 (개체 포함)
  /export/vendors    거래처
  /export/workers    근무자
"""
import io
from datetime import date, datetime
from fastapi import APIRouter, Request, HTTPException
from fastapi.responses import StreamingResponse
from sqlmodel import Session, select

from database import (
    engine, User, Project, Vendor, Expense, Quote, QuoteItem,
    Item, Equipment, EquipmentUnit, Worker, Attendance,
)

router = APIRouter()

HEADER_BG = "0A0E1A"


def _user(request: Request):
    uid = request.session.get("user_id")
    if not uid:
        raise HTTPException(303, headers={"Location": "/login"})
    with Session(engine, expire_on_commit=False) as s:
        return s.get(User, uid)


def _st(request: Request, name: str) -> bool:
    try:
        return bool(getattr(request.state, name, False))
    except Exception:
        return False


def _can_view_amounts(request: Request) -> bool:
    try:
        return bool(getattr(request.state, "can_view_amounts", True))
    except Exception:
        return True


def _new_book():
    from openpyxl import Workbook
    return Workbook()


def _style_header(ws):
    from openpyxl.styles import Font, PatternFill, Alignment
    font = Font(bold=True, color="FFFFFF")
    fill = PatternFill("solid", fgColor=HEADER_BG)
    center = Alignment(horizontal="center", vertical="center")
    for cell in ws[1]:
        cell.font = font
        cell.fill = fill
        cell.alignment = center
    ws.freeze_panes = "A2"


def _autofit(ws, max_width=55):
    for col in ws.columns:
        longest = 0
        letter = None
        for cell in col:
            if letter is None:
                letter = cell.column_letter
            v = cell.value
            if v is None:
                continue
            for line in str(v).split("\n"):
                # 한글은 폭을 넓게 계산
                w = sum(2 if ord(ch) > 0x2000 else 1 for ch in line)
                longest = max(longest, w)
        if letter:
            ws.column_dimensions[letter].width = min(max(10, longest + 2), max_width)


def _send(wb, filename_base: str):
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    stamp = datetime.now().strftime("%Y%m%d_%H%M")
    fname = f"{filename_base}_{stamp}.xlsx"
    from urllib.parse import quote
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={
            "Content-Disposition": (
                f"attachment; filename=export_{stamp}.xlsx; "
                f"filename*=UTF-8''{quote(fname)}"
            )
        },
    )


def _d(v):
    """날짜 → 문자열 (빈 값 안전)"""
    if not v:
        return ""
    try:
        return v.isoformat()
    except Exception:
        return str(v)


# ============================================================
# 프로젝트
# ============================================================

@router.get("/export/projects")
def export_projects(request: Request):
    _user(request)
    show_money = _can_view_amounts(request)
    wb = _new_book()
    ws = wb.active
    ws.title = "프로젝트"
    sv_settle, sv_sub, sv_wage, sv_profit = (_st(request, "can_view_settle"), _st(request, "can_view_sub"),
                                             _st(request, "can_view_wage"), _st(request, "can_view_profit"))
    headers = ["코드", "프로젝트명", "거래처", "행사일", "장소", "카테고리", "진행상태"]
    if sv_settle:
        headers += ["정산상태"]
    if show_money:
        headers += ["공급가", "부가세", "매출(총액)"]
    if sv_sub:
        headers += ["외주 지급"]
    if sv_wage:
        headers += ["알바 급여"]
    if sv_profit:
        headers += ["순이익"]
    if sv_settle:
        headers += ["입금일", "결제수단"]
    headers += ["특이사항", "메모"]
    ws.append(headers)
    _style_header(ws)

    with Session(engine, expire_on_commit=False) as s:
        vendors = {v.id: v.name for v in s.exec(select(Vendor)).all()}
        for p in s.exec(select(Project).order_by(Project.created_at.desc())).all():
            exp = sum(e.amount for e in s.exec(
                select(Expense).where(Expense.project_id == p.id)).all())
            wage = sum(a.total_wage for a in s.exec(
                select(Attendance).where(Attendance.project_id == p.id)).all())
            supply = p.supply_amount or p.revenue
            row = [
                p.code or "", p.name or "",
                vendors.get(p.vendor_id, "") if p.vendor_id else "",
                _d(p.event_date), p.location or "", p.categories or "",
                p.status or "",
            ]
            if sv_settle:
                row += [p.settlement_status or ""]
            if show_money:
                row += [supply, p.vat_amount or 0, p.revenue or 0]
            if sv_sub:
                row += [exp]
            if sv_wage:
                row += [wage]
            if sv_profit:
                row += [supply - exp - wage]
            if sv_settle:
                row += [_d(p.paid_date), p.payment_method or ""]
            row += [p.special_notes or "", p.memo or ""]
            ws.append(row)
    _autofit(ws)
    return _send(wb, "프로젝트")


# ============================================================
# 견적서
# ============================================================

@router.get("/export/quotes")
def export_quotes(request: Request):
    _user(request)
    show_money = _can_view_amounts(request)
    wb = _new_book()

    ws = wb.active
    ws.title = "견적서"
    headers = ["견적번호", "견적일", "유형", "거래처", "프로젝트명", "행사일", "상태"]
    if show_money:
        headers += ["소계", "할인", "공급가", "부가세", "총액"]
    headers += ["부가세방식", "유효기간", "결제조건", "납품조건", "특이사항"]
    ws.append(headers)
    _style_header(ws)

    ws2 = wb.create_sheet("견적품목")
    h2 = ["견적번호", "순번", "품명", "규격", "수량", "단위"]
    if show_money:
        h2 += ["원가", "단가", "금액"]
    h2 += ["비고"]
    ws2.append(h2)
    _style_header(ws2)

    with Session(engine, expire_on_commit=False) as s:
        for q in s.exec(select(Quote).order_by(Quote.quote_date.desc())).all():
            row = [
                q.quote_number or "", _d(q.quote_date), q.quote_type or "",
                q.vendor_name or "", q.project_name or "",
                _d(q.event_date), q.status or "",
            ]
            if show_money:
                row += [q.subtotal or 0, q.discount or 0,
                        q.after_discount or 0, q.vat or 0, q.total or 0]
            row += [q.vat_mode or "", _d(q.valid_until),
                    q.payment_terms or "", q.delivery_terms or "", q.notes or ""]
            ws.append(row)

            items = s.exec(select(QuoteItem)
                           .where(QuoteItem.quote_id == q.id)
                           .order_by(QuoteItem.seq)).all()
            for it in items:
                r2 = [q.quote_number or "", it.seq, it.name or "",
                      it.spec or "", it.quantity or 0, it.unit or ""]
                if show_money:
                    r2 += [getattr(it, "cost_price", 0) or 0,
                           it.unit_price or 0, it.amount or 0]
                r2 += [it.memo or ""]
                ws2.append(r2)
    _autofit(ws)
    _autofit(ws2)
    return _send(wb, "견적서")


# ============================================================
# 하청/외주 지급
# ============================================================

@router.get("/export/expenses")
def export_expenses(request: Request, pay: str = ""):
    """pay=미지급 / 지급완료 로 필터 가능"""
    _user(request)
    show_money = _st(request, "can_view_sub")
    wb = _new_book()
    ws = wb.active
    ws.title = "하청외주지급"
    headers = ["발주일", "프로젝트", "거래처", "역할", "작업내용", "상세사양"]
    if show_money:
        headers += ["공급가", "부가세", "지급액(총액)"]
    headers += ["지급상태", "지급예정일", "지급일", "결제수단", "세금계산서", "지급메모"]
    ws.append(headers)
    _style_header(ws)

    with Session(engine, expire_on_commit=False) as s:
        projects = {p.id: p.name for p in s.exec(select(Project)).all()}
        vendors = {v.id: v.name for v in s.exec(select(Vendor)).all()}
        rows = s.exec(select(Expense).order_by(Expense.expense_date.desc())).all()
        pay_f = (pay or "").strip()
        if pay_f in ("미지급", "지급완료"):
            rows = [e for e in rows
                    if (getattr(e, "pay_status", "미지급") or "미지급") == pay_f]
        for e in rows:
            row = [
                _d(e.expense_date),
                projects.get(e.project_id, ""),
                vendors.get(e.vendor_id, "") if e.vendor_id else (e.vendor_name or ""),
                e.category or "", e.description or "",
                getattr(e, "spec_detail", "") or "",
            ]
            if show_money:
                row += [getattr(e, "supply_amount", 0) or 0,
                        getattr(e, "vat_amount", 0) or 0, e.amount or 0]
            row += [
                getattr(e, "pay_status", "미지급") or "미지급",
                _d(getattr(e, "pay_due_date", None)),
                _d(getattr(e, "pay_date", None)),
                e.payment_method or "",
                "발행" if getattr(e, "has_tax_invoice", False) else "",
                getattr(e, "pay_memo", "") or "",
            ]
            ws.append(row)
    _autofit(ws)
    return _send(wb, "하청외주지급")


# ============================================================
# 견적 단가표
# ============================================================

@router.get("/export/items")
def export_items(request: Request):
    _user(request)
    show_money = _can_view_amounts(request)
    wb = _new_book()
    ws = wb.active
    ws.title = "견적단가표"
    headers = ["코드", "품명", "규격", "단위", "카테고리"]
    if show_money:
        headers += ["수입단가", "대리점가", "납품단가", "렌탈1일", "렌탈보증금"]
    headers += ["메모", "사용여부"]
    ws.append(headers)
    _style_header(ws)

    with Session(engine, expire_on_commit=False) as s:
        rows = s.exec(select(Item).where(Item.usage_type == "quote")
                      .order_by(Item.code)).all()
        for i in rows:
            row = [i.code or "", i.name or "", i.spec or "",
                   i.unit or "", i.category or ""]
            if show_money:
                row += [i.import_price or 0, i.dealer_price or 0,
                        i.consumer_price or 0, i.rental_daily or 0,
                        i.rental_deposit or 0]
            row += [i.memo or "", "사용" if i.is_active else "미사용"]
            ws.append(row)
    _autofit(ws)
    return _send(wb, "견적단가표")


# ============================================================
# 재고 / 장비
# ============================================================

@router.get("/export/inventory")
def export_inventory(request: Request):
    _user(request)
    show_money = _can_view_amounts(request)
    wb = _new_book()

    ws = wb.active
    ws.title = "장비종류"
    ws.append(["장비명", "모델", "제조사", "카테고리", "규격", "보유수량", "비고"])
    _style_header(ws)

    ws2 = wb.create_sheet("장비개체")
    h2 = ["관리코드", "장비명", "모델", "시리얼", "상태", "보관위치",
          "관리부서", "담당자", "구입처", "구입일"]
    if show_money:
        h2 += ["구입금액", "렌탈1일"]
    h2 += ["최근점검일", "메모"]
    ws2.append(h2)
    _style_header(ws2)

    with Session(engine, expire_on_commit=False) as s:
        for eq in s.exec(select(Equipment).order_by(Equipment.category, Equipment.name)).all():
            units = s.exec(select(EquipmentUnit)
                           .where(EquipmentUnit.equipment_id == eq.id)
                           .order_by(EquipmentUnit.asset_code)).all()
            ws.append([eq.name or "", eq.model or "", eq.manufacturer or "",
                       eq.category or "", eq.spec or "", len(units), eq.memo or ""])
            for u in units:
                r2 = [u.asset_code or "", eq.name or "", eq.model or "",
                      u.serial_number or "", u.status or "", u.location or "",
                      u.department or "", u.manager or "",
                      u.purchase_vendor or "", _d(u.purchase_date)]
                if show_money:
                    r2 += [u.purchase_price or 0, u.rental_price_daily or 0]
                r2 += [_d(u.last_check_date), u.memo or ""]
                ws2.append(r2)
    _autofit(ws)
    _autofit(ws2)
    return _send(wb, "재고장비")


# ============================================================
# 거래처
# ============================================================

@router.get("/export/vendors")
def export_vendors(request: Request):
    _user(request)
    show_money = _can_view_amounts(request)
    wb = _new_book()
    ws = wb.active
    ws.title = "거래처"
    headers = ["거래처명", "담당자", "전화", "이메일", "사업자번호", "구분", "주소", "메모"]
    if show_money:
        headers += ["프로젝트수", "총매출", "하청지급액", "미지급액"]
    ws.append(headers)
    _style_header(ws)

    with Session(engine, expire_on_commit=False) as s:
        all_proj = s.exec(select(Project)).all()
        all_exp = s.exec(select(Expense)).all()
        for v in s.exec(select(Vendor).order_by(Vendor.name)).all():
            row = [v.name or "", v.contact_person or "", v.phone or "",
                   v.email or "", v.biz_number or "",
                   "대리점" if v.vendor_type == "dealer" else "일반",
                   v.address or "", v.memo or ""]
            if show_money:
                ps = [p for p in all_proj if p.vendor_id == v.id]
                es = [e for e in all_exp if e.vendor_id == v.id]
                unpaid = sum(e.amount for e in es
                             if (getattr(e, "pay_status", "미지급") or "미지급") == "미지급")
                row += [len(ps), sum(p.revenue or 0 for p in ps),
                        sum(e.amount or 0 for e in es), unpaid]
            ws.append(row)
    _autofit(ws)
    return _send(wb, "거래처")


# ============================================================
# 근무자
# ============================================================

@router.get("/export/workers")
def export_workers(request: Request):
    user = _user(request)
    from permissions import is_admin
    if not is_admin(user):
        raise HTTPException(403, "근무자 정보 내보내기는 관리자만 가능합니다.")
    show_money = _can_view_amounts(request)
    wb = _new_book()
    ws = wb.active
    ws.title = "근무자"
    headers = ["이름", "구분", "직책", "연락처", "은행", "계좌"]
    if show_money:
        headers += ["기본일당", "월급"]
    headers += ["입사일", "퇴사일", "재직상태", "메모"]
    ws.append(headers)
    _style_header(ws)

    with Session(engine, expire_on_commit=False) as s:
        for w in s.exec(select(Worker).order_by(Worker.name)).all():
            row = [w.name or "", w.employee_type or "", w.position or "",
                   w.phone or "", w.bank or "", w.account or ""]
            if show_money:
                row += [w.default_daily_wage or 0, w.monthly_salary or 0]
            row += [_d(w.hire_date), _d(w.resign_date),
                    w.employment_status or "", w.memo or ""]
            ws.append(row)
    _autofit(ws)
    return _send(wb, "근무자")
